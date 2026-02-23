import io
from PIL import Image
from fpdf import FPDF
from docx import Document
import pandas as pd
import os

class Converter:

    @staticmethod
    def text_to_pdf(text: str, filename: str = "output.pdf") -> io.BytesIO:
        pdf = FPDF()
        pdf.add_page()

        try:
            pdf.add_font('DejaVu', '', 'src/services/DejaVu Sans/DejaVuSans-Oblique.ttf', uni=True)
            pdf.set_font('DejaVu', size=12)
        except Exception as e:
            print("font err", e)
            pdf.set_font('Arial', size=12)

        for line in text.split('\n'):
            pdf.cell(0, 10, str(line), ln=True)

        pdf_bytes = io.BytesIO()
        pdf_bytes.write(pdf.output(dest='S').encode('latin1'))
        pdf_bytes.seek(0)
        
        return pdf_bytes
    

    @staticmethod
    def text_to_docx(text: str, filename: str = "output.docx") -> io.BytesIO:
        doc = Document()
        for line in text.split('\n'):
            doc.add_paragraph(line)

        docx_bytes = io.BytesIO()
        doc.save(docx_bytes)
        docx_bytes.seek(0)
        return docx_bytes
    
    @staticmethod
    def img_conv(file: io.BytesIO, need_type: str) -> io.BytesIO:
        with Image.open(file) as i:
            out = io.BytesIO()
            i.save(out, format=need_type.upper())
            out.seek(0)

        return out